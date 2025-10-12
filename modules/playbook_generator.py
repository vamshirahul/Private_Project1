import io, zipfile, math
from pathlib import Path
import streamlit as st
from openai import OpenAI

# === Utils from your toolkit ===
from utils.prompts import PLAYBOOK_SYSTEM, PLAYBOOK_USER, PLAYBOOK_DETAILED_USER
from utils.helpers import (
    text_to_bytes, get_api_key, new_job_id, save_artifact, record_dataset
)

# === Optional backend client (RAG + LLM in one) ===
import requests
def ask_backend(question: str, tags=None, k=6, base_url="http://localhost:8000",
                model="gpt-4o-mini", temperature=0.2, metadata=None):
    payload = {
        "question": question, "tags": tags or [], "k": k,
        "model": model, "temperature": float(temperature),
        "metadata": metadata or {},
    }
    r = requests.post(f"{base_url}/ask", json=payload, timeout=120)
    r.raise_for_status()
    return r.json()

# === Simple DOCX maker (kept local to avoid extra imports elsewhere) ===
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING
def md_to_docx_bytes(md_text: str, title="Integration Playbook") -> bytes:
    doc = Document()
    s = doc.styles['Normal']; s.font.name='Calibri'; s.font.size=Pt(10.5)
    doc.add_heading(title, 0)
    for line in md_text.splitlines():
        if line.startswith("### "): doc.add_heading(line[4:].strip(), 2)
        elif line.startswith("## "): doc.add_heading(line[3:].strip(), 1)
        elif line.startswith("# "): doc.add_heading(line[2:].strip(), 0)
        else:
            p = doc.add_paragraph(line); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

# === ZIP pack helper ===
def make_export_zip(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()

def _get_client():
    api_key = get_api_key(st)
    if not api_key:
        st.error("Missing OPENAI_API_KEY (secrets or env).")
        return None
    return OpenAI(api_key=api_key)

def _clean_app_list(raw: str):
    return [a.strip() for a in raw.replace("\r", "\n").replace(",", "\n").split("\n") if a.strip()]

def render():
    st.header("🧩 Integration Playbook Generator")
    st.caption("Standard summary mode + Advanced per-application mode")

    tab_std, tab_adv = st.tabs(["⭐ Standard Mode", "🛠 Advanced Mode"])

    # -------------------------
    # ⭐ STANDARD MODE
    # -------------------------
    with tab_std:
        with st.expander("⚙️ Settings", expanded=False):
            colm1, colm2 = st.columns(2)
            with colm1:
                model_std = st.selectbox("Model", ["gpt-4o-mini", "gpt-4o"], index=0, key="mdl_std")
            with colm2:
                temp_std = st.slider("Creativity", 0.0, 1.0, 0.3, 0.1, key="tmp_std")

        with st.form("pg_form_standard"):
            company_name = st.text_input("Company Name")
            company_size = st.selectbox("Company Size", ["Small", "Mid-size", "Large"])
            key_systems = st.text_area("Key Systems (comma or newline separated)")
            day1_priority = st.text_area("Day-1 Priorities")
            integration_approach = st.selectbox("Integration Approach", ["Full Integration", "Phased", "Hybrid"])
            timeline = st.text_input("Integration Timeline (e.g., 30/60/90 or 100 days)")
            use_backend = st.checkbox("Use backend /ask (RAG + logging)", value=False, key="std_use_backend")
            base_url = st.text_input("Backend URL", value="http://localhost:8000", key="std_base") if use_backend else ""
            submitted = st.form_submit_button("Generate Playbook")

        if submitted:
            apps = _clean_app_list(key_systems)
            user_prompt = PLAYBOOK_USER.format(
                company_name=company_name.strip(),
                company_size=company_size,
                key_systems=", ".join(apps),
                day1_priority=day1_priority.strip(),
                integration_approach=integration_approach,
                timeline=timeline.strip()
            )

            job_id = new_job_id("playbook_std")
            with st.spinner("Generating playbook..."):
                if use_backend:
                    q = ("Create a concise integration playbook using the structure in the prompt. "
                         "Return markdown with headings, tables where helpful.\n\n" + user_prompt)
                    res = ask_backend(q, tags=["Integration","Playbook"], k=6,
                                      base_url=base_url, model=model_std, temperature=temp_std,
                                      metadata={"module":"playbook_standard"})
                    md = res["answer"]; citations = res.get("citations", [])
                else:
                    client = _get_client()
                    if not client: return
                    resp = client.chat.completions.create(
                        model=model_std,
                        messages=[{"role": "system", "content": PLAYBOOK_SYSTEM},
                                  {"role": "user", "content": user_prompt}],
                        temperature=temp_std,
                    )
                    md = resp.choices[0].message.content
                    citations = []

            # KPIs
            c1,c2,c3 = st.columns(3)
            with c1: st.markdown(f'<div class="kpi">🧠 Model<br><b>{model_std}</b></div>', unsafe_allow_html=True)
            with c2: st.markdown(f'<div class="kpi">🎛 Temperature<br><b>{temp_std}</b></div>', unsafe_allow_html=True)
            with c3: st.markdown(f'<div class="kpi">💾 Logged<br><b>{"Yes" if st.session_state.get("capture_outputs") else "No"}</b></div>', unsafe_allow_html=True)

            # Output card
            st.markdown('<div class="card">', unsafe_allow_html=True)
            st.markdown("#### ✅ Output")
            st.markdown(md)
            st.markdown('</div>', unsafe_allow_html=True)

            # Citations (if backend used)
            if citations:
                with st.expander("🔎 Sources used"):
                    for i, c in enumerate(citations, start=1):
                        st.markdown(f"**CTX {i}** · score={c['score']:.3f}")
                        st.write(c["text"])
                        st.markdown("---")

            # Exports
            md_bytes = text_to_bytes(md)
            st.download_button("⬇️ Download Playbook (.md)", data=md_bytes,
                               file_name="integration_playbook.md", mime="text/markdown", use_container_width=True)

            docx_bytes = md_to_docx_bytes(md, title=f"{company_name} - Integration Playbook")
            st.download_button("📄 Download DOCX", data=docx_bytes,
                               file_name="integration_playbook.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)

            export_zip = make_export_zip({
                "playbook.md": md_bytes,
                "playbook.docx": docx_bytes,
                "toc.txt": text_to_bytes(", ".join(apps))
            })
            st.download_button("📦 Download Export Pack (.zip)", data=export_zip,
                               file_name="integration_export_pack.zip", mime="application/zip",
                               use_container_width=True)

            # Artifacts + dataset
            path_md = save_artifact(job_id, "integration_playbook.md", md_bytes)
            if st.session_state.get("capture_outputs", False):
                record_dataset(
                    module="playbook_standard",
                    job_id=job_id,
                    inputs={"company_name":company_name,"company_size":company_size,"apps":apps,
                            "day1_priority":day1_priority,"integration_approach":integration_approach,
                            "timeline":timeline,"model":model_std,"temperature":temp_std,
                            "use_backend":use_backend},
                    output_text=md,
                    context_text=None,
                    artifacts={"playbook_md": path_md},
                    quality="unreviewed"
                )

    # -------------------------
    # 🛠 ADVANCED MODE (per-app)
    # -------------------------
    with tab_adv:
        with st.expander("⚙️ Settings", expanded=False):
            c1, c2, c3, c4 = st.columns(4)
            with c1: model_name = st.selectbox("Model", ["gpt-4o", "gpt-4o-mini"], index=0, key="mdl_adv")
            with c2: temperature = st.slider("Creativity", 0.0, 1.0, 0.3, 0.1, key="tmp_adv")
            with c3: batch_size = st.number_input("Batch size", 10, 50, 25, 5, key="bsz_adv")
            with c4:
                use_backend_adv = st.checkbox("Use backend /ask (RAG)", value=True, key="adv_use_backend")
                base_url_adv = st.text_input("Backend URL", value="http://localhost:8000", key="adv_base") if use_backend_adv else ""

        with st.form("pg_form_advanced"):
            company_name_a = st.text_input("Company Name", key="adv_name")
            company_size_a = st.selectbox("Company Size", ["Small", "Mid-size", "Large"], key="adv_size")
            apps_text = st.text_area(
                "Applications list (comma or newline separated)",
                placeholder="SAP ECC\nSalesforce\nWorkday\nServiceNow\nJD Edwards\n...",
                height=200, key="adv_apps")
            integration_approach_a = st.selectbox("Integration Approach", ["Full Integration", "Phased", "Hybrid"], key="adv_appr")
            timeline_a = st.text_input("Overall Integration Timeline", value="100 days", key="adv_time")
            submitted_a = st.form_submit_button("Generate Detailed Per-App Playbook")

        if submitted_a:
            apps = _clean_app_list(apps_text)
            if not apps:
                st.error("Please provide at least one application.")
                return

            job_id = new_job_id("playbook_adv")
            batches = [apps[i:i+int(batch_size)] for i in range(0, len(apps), int(batch_size))]
            progress = st.progress(0, text=f"Generating {len(apps)} apps in {len(batches)} batch(es)...")
            results = []
            all_citations = []

            for idx, group in enumerate(batches, start=1):
                apps_block = "\n".join(group)
                base_prompt = PLAYBOOK_DETAILED_USER.format(
                    applications=apps_block,
                    company_name=company_name_a.strip(),
                    company_size=company_size_a,
                    integration_approach=integration_approach_a,
                    timeline=timeline_a.strip(),
                    app_name="{app_name}"
                )

                with st.spinner(f"Batch {idx}/{len(batches)}…"):
                    if use_backend_adv:
                        q = ("Create detailed per-application integration plans for these apps, "
                             "including Overview, Strategy, Data/Cutover, Risks, Wave; end with a summary table.\n\n" + base_prompt)
                        res = ask_backend(q, tags=["Integration","Cutover"], k=6,
                                          base_url=base_url_adv, model=model_name, temperature=temperature,
                                          metadata={"module":"playbook_advanced","batch_index": idx})
                        batch_md = res["answer"]; citations = res.get("citations", [])
                        all_citations.extend(citations)
                    else:
                        client = _get_client()
                        if not client: return
                        resp = client.chat.completions.create(
                            model=model_name,
                            messages=[{"role": "system", "content": PLAYBOOK_SYSTEM},
                                      {"role": "user", "content": base_prompt}],
                            temperature=temperature,
                        )
                        batch_md = resp.choices[0].message.content
                results.append(batch_md)
                progress.progress(idx/len(batches), text=f"Completed {idx}/{len(batches)}")

            final_md = "\n\n---\n\n".join(results)

            # KPI row
            k1,k2,k3 = st.columns(3)
            with k1: st.markdown(f'<div class="kpi">🔁 Batches<br><b>{len(batches)}</b></div>', unsafe_allow_html=True)
            with k2: st.markdown(f'<div class="kpi">🧠 Model<br><b>{model_name}</b></div>', unsafe_allow_html=True)
            with k3: st.markdown(f'<div class="kpi">💾 Logged<br><b>{"Yes" if st.session_state.get("capture_outputs") else "No"}</b></div>', unsafe_allow_html=True)

            # Output card
            st.markdown('<div class="card">', unsafe_allow_html=True)
            st.markdown("### Detailed Plan")
            st.markdown(final_md)
            st.markdown('</div>', unsafe_allow_html=True)

            # Citations summary
            if use_backend_adv and all_citations:
                with st.expander("🔎 Sources used"):
                    for i, c in enumerate(all_citations, start=1):
                        st.markdown(f"**CTX {i}** · score={c['score']:.3f}")
                        st.write(c["text"])
                        st.markdown("---")

            # Exports
            md_bytes = text_to_bytes(final_md)
            st.download_button("⬇️ Download Detailed Playbook (.md)", data=md_bytes,
                               file_name="integration_playbook_detailed.md", mime="text/markdown", use_container_width=True)

            docx_bytes = md_to_docx_bytes(final_md, title=f"{company_name_a} - Integration Playbook (Detailed)")
            st.download_button("📄 Download DOCX", data=docx_bytes,
                               file_name="integration_playbook_detailed.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)

            export_zip = make_export_zip({
                "playbook_detailed.md": md_bytes,
                "playbook_detailed.docx": docx_bytes,
                "toc.txt": text_to_bytes("\n".join(apps)),
            })
            st.download_button("📦 Download Export Pack (.zip)", data=export_zip,
                               file_name="integration_export_pack.zip", mime="application/zip", use_container_width=True)

            # Artifacts + dataset
            path_md = save_artifact(job_id, "integration_playbook_detailed.md", md_bytes)
            if st.session_state.get("capture_outputs", False):
                record_dataset(
                    module="playbook_advanced",
                    job_id=job_id,
                    inputs={
                        "company_name": company_name_a, "company_size": company_size_a,
                        "apps": apps, "integration_approach": integration_approach_a,
                        "timeline": timeline_a, "batch_size": int(batch_size),
                        "model": model_name, "temperature": temperature,
                        "use_backend": use_backend_adv
                    },
                    output_text=final_md,
                    context_text=None,  # if you want, store concatenated CTX here
                    artifacts={"playbook_md": path_md},
                    quality="unreviewed"
                )
